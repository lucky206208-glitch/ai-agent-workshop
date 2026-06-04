from __future__ import annotations

import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "decks" / "02-github-pr-control-bento-swiss.pptx"
LINKS = ROOT / "docs" / "training-pr-links.json"
ASSETS = ROOT / "decks" / "assets" / "lecture-02"

FONT_KO = "Malgun Gothic"
FONT_MONO = "Consolas"

INK = RGBColor(17, 17, 17)
MUTED = RGBColor(68, 68, 68)
PAPER = RGBColor(248, 248, 242)
ACCENT = RGBColor(232, 0, 13)
NAVY = RGBColor(26, 26, 46)
TEAL = RGBColor(78, 205, 196)
YELLOW = RGBColor(232, 255, 59)
CORAL = RGBColor(255, 107, 107)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(221, 221, 221)
GREEN = RGBColor(31, 136, 61)
RED = RGBColor(207, 34, 46)
TRAINING_MARKER = "_".join(["TRAINING", "SECRET", "DO", "NOT", "USE"])


def inches(value: float):
    return Inches(value)


def set_fill(slide, color=PAPER) -> None:
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 24,
    bold: bool = False,
    color=INK,
    font: str = FONT_KO,
    align=None,
    valign=None,
    name: str | None = None,
):
    box = slide.shapes.add_textbox(inches(x), inches(y), inches(w), inches(h))
    if name:
        box.name = name
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = inches(0.03)
    frame.margin_right = inches(0.03)
    frame.margin_top = inches(0.02)
    frame.margin_bottom = inches(0.02)
    if valign:
        frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_link_text(slide, text: str, url: str, x: float, y: float, w: float, h: float):
    box = slide.shapes.add_textbox(inches(x), inches(y), inches(w), inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = FONT_MONO
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(9, 105, 218)
    run.hyperlink.address = url
    return box


def add_link_button(slide, label: str, url: str, x: float, y: float, w: float = 2.85, h: float = 0.38):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inches(x), inches(y), inches(w), inches(h))
    shape.adjustments[0] = 0.18
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(9, 105, 218)
    box = add_text(slide, label, x + 0.12, y + 0.08, w - 0.24, 0.17, size=9, bold=True, color=RGBColor(9, 105, 218), font=FONT_KO)
    run = box.text_frame.paragraphs[0].runs[0]
    run.hyperlink.address = url
    return shape


def add_rule(slide, x: float, y: float, w: float, h: float = 0.025, color=ACCENT):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inches(x), inches(y), inches(w), inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    return shape


def add_section_label(slide, label: str) -> None:
    marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inches(0.48), inches(0.38), inches(0.14), inches(0.14))
    marker.name = "kicker-marker"
    marker.fill.solid()
    marker.fill.fore_color.rgb = ACCENT
    marker.line.color.rgb = ACCENT
    add_text(
        slide,
        label,
        0.74,
        0.31,
        4.6,
        0.28,
        size=10,
        bold=True,
        color=ACCENT,
        font=FONT_MONO,
        valign=MSO_ANCHOR.MIDDLE,
        name="kicker-label",
    )
    add_rule(slide, 0.48, 0.9, 2.85)


def add_footer(slide, number: int) -> None:
    add_text(
        slide,
        f"{number:02d} / Design reference: corazzon/pptx-design-styles (MIT)",
        9.45,
        7.1,
        3.25,
        0.18,
        size=7,
        color=RGBColor(115, 115, 115),
        font=FONT_MONO,
        align=PP_ALIGN.RIGHT,
    )


def add_title(slide, label: str, title: str, subtitle: str, number: int) -> None:
    add_section_label(slide, label)
    add_text(slide, title, 0.72, 1.06, 10.05, 1.16, size=28, bold=True)
    add_text(slide, subtitle, 0.72, 2.34, 9.65, 0.72, size=13, color=MUTED)
    add_footer(slide, number)


def add_card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    title: str,
    body: str,
    *,
    fill=WHITE,
    dark=False,
    body_size: int = 9,
):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inches(x), inches(y), inches(w), inches(h))
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = LINE if not dark else fill
    title_color = WHITE if dark else INK
    body_color = RGBColor(230, 230, 230) if dark else MUTED
    add_text(slide, title, x + 0.18, y + 0.16, w - 0.36, 0.34, size=13, bold=True, color=title_color)
    add_text(slide, body, x + 0.18, y + 0.58, w - 0.36, h - 0.72, size=body_size, color=body_color)
    return shape


def add_status_chip(slide, x: float, y: float, label: str, color, text_color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inches(x), inches(y), inches(1.65), inches(0.33))
    shape.adjustments[0] = 0.5
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    add_text(slide, label, x + 0.04, y + 0.06, 1.55, 0.16, size=9, bold=True, color=text_color, align=PP_ALIGN.CENTER)


def add_asset(slide, filename: str, x: float = 5.55, y: float = 2.76, w: float = 7.15, h: float = 4.02) -> None:
    path = ASSETS / filename
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inches(x), inches(y), inches(w), inches(h))
    frame.fill.solid()
    frame.fill.fore_color.rgb = WHITE
    frame.line.color.rgb = LINE
    if not path.exists():
        add_card(slide, x + 0.2, y + 0.2, w - 0.4, h - 0.4, "Screenshot needed", filename)
        return
    with Image.open(path) as img:
        img_w, img_h = img.size
    ratio = min(w / img_w, h / img_h)
    pic_w = img_w * ratio
    pic_h = img_h * ratio
    pic_x = x + (w - pic_w) / 2
    pic_y = y + (h - pic_h) / 2
    slide.shapes.add_picture(str(path), inches(pic_x), inches(pic_y), width=inches(pic_w), height=inches(pic_h))


def add_step_flow(slide, y: float, items: list[tuple[str, str, object]]) -> None:
    x = 0.72
    card_w = 3.55 if len(items) <= 3 else 2.72
    step_gap = 4.1 if len(items) <= 3 else 3.05
    body_size = 9 if len(items) <= 3 else 8
    add_rule(slide, x + 0.22, y - 0.18, step_gap * (len(items) - 1) + card_w - 0.44, 0.018, RGBColor(198, 198, 188))
    for idx, (title, body, fill) in enumerate(items):
        step_x = x + idx * step_gap
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, inches(step_x + 0.15), inches(y - 0.36), inches(0.34), inches(0.34))
        circle.fill.solid()
        circle.fill.fore_color.rgb = ACCENT if idx == 0 else RGBColor(245, 245, 238)
        circle.line.color.rgb = ACCENT
        add_text(slide, str(idx + 1), step_x + 0.15, y - 0.285, 0.34, 0.12, size=8, bold=True, color=WHITE if idx == 0 else ACCENT, align=PP_ALIGN.CENTER, font=FONT_MONO)
        add_card(slide, step_x, y, card_w, 1.42, title, body, fill=fill, dark=fill == NAVY, body_size=body_size)
        if idx < len(items) - 1:
            chevron = slide.shapes.add_shape(
                MSO_SHAPE.CHEVRON,
                inches(step_x + card_w + 0.13),
                inches(y + 0.51),
                inches(0.22),
                inches(0.28),
            )
            chevron.fill.solid()
            chevron.fill.fore_color.rgb = ACCENT
            chevron.line.color.rgb = ACCENT


def add_bullet_stack(slide, x: float, y: float, bullets: list[str], color=INK):
    for idx, bullet in enumerate(bullets):
        add_rule(slide, x, y + idx * 0.47 + 0.09, 0.12, 0.12, ACCENT)
        add_text(slide, bullet, x + 0.24, y + idx * 0.47, 5.55, 0.34, size=11, color=color)


def add_glossary_showcase(slide, items: list[tuple[str, str, str, object]]) -> None:
    for idx, (term, body, example, accent) in enumerate(items):
        y = 3.02 + idx * 1.23
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inches(0.72), inches(y), inches(11.8), inches(1.04))
        shape.adjustments[0] = 0.04
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = LINE
        rail = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inches(0.72), inches(y), inches(0.16), inches(1.04))
        rail.fill.solid()
        rail.fill.fore_color.rgb = accent
        rail.line.color.rgb = accent
        chip = slide.shapes.add_shape(MSO_SHAPE.OVAL, inches(1.05), inches(y + 0.29), inches(0.46), inches(0.46))
        chip.fill.solid()
        chip.fill.fore_color.rgb = accent
        chip.line.color.rgb = accent
        add_text(slide, str(idx + 1), 1.05, y + 0.39, 0.46, 0.15, size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER, font=FONT_MONO)
        add_text(slide, term, 1.72, y + 0.18, 2.7, 0.3, size=14, bold=True, color=INK)
        add_text(slide, body, 4.35, y + 0.17, 4.15, 0.35, size=11, color=INK)
        add_text(slide, example, 8.72, y + 0.17, 3.38, 0.48, size=9, color=MUTED)


def add_command_box(slide, x: float, y: float, w: float, h: float, title: str, lines: list[str], *, fill=WHITE, dark=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inches(x), inches(y), inches(w), inches(h))
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill if dark else LINE
    title_color = WHITE if dark else INK
    body_color = RGBColor(232, 232, 232) if dark else MUTED
    add_text(slide, title, x + 0.18, y + 0.16, w - 0.36, 0.28, size=13, bold=True, color=title_color)
    add_text(slide, "\n".join(lines), x + 0.18, y + 0.58, w - 0.36, h - 0.72, size=9, color=body_color, font=FONT_MONO)
    return shape


def create_deck() -> None:
    prs = Presentation()
    prs.slide_width = inches(13.333)
    prs.slide_height = inches(7.5)
    blank = prs.slide_layouts[6]
    links = json.loads(LINKS.read_text(encoding="utf-8"))
    scenarios = {item["id"]: item for item in links["scenarios"]}

    slide_no = 1

    slide = prs.slides.add_slide(blank)
    set_fill(slide)
    add_section_label(slide, "LECTURE 02 / GITHUB PR CONTROL")
    add_text(slide, "AI agent PR을\nmerge해도 되는지 판단하기", 0.72, 1.25, 7.2, 1.55, size=36, bold=True)
    add_text(slide, "Git 명령어 암기가 아니라 PR 화면에서 변경 범위, 실제 diff, 자동검사 결과를 읽는 수업입니다.", 0.76, 2.9, 7.6, 0.58, size=14, color=MUTED)
    add_card(slide, 0.72, 4.08, 3.4, 1.35, "핵심 질문", "이 변경을 main에 합쳐도 되는가? 근거는 파일, diff, checks에서 찾습니다.", fill=YELLOW)
    add_card(slide, 4.38, 3.72, 3.0, 1.72, "읽는 순서", "Files changed -> Diff -> Checks -> Decision 순서로 같은 PR을 반복해서 봅니다.", fill=TEAL)
    add_card(slide, 7.68, 4.08, 3.35, 1.35, "결정", "Merge / Request changes / Hold 중 하나를 고르고 이유를 말합니다.", fill=NAVY, dark=True)
    add_rule(slide, 11.55, 1.15, 0.25, 4.6)
    add_footer(slide, slide_no)
    slide_no += 1

    specs = [
        {
            "label": "PROBLEM",
            "title": "속도보다 merge 판단 기준이 중요하다",
            "subtitle": "예: 제목 한 줄 수정 요청이 왔는데 agent가 스타일, 문서, 설정까지 바꾸면 속도보다 검토 기준이 먼저입니다.",
            "cards": [
                ("빠른 생성", "요청이 작아도 agent는 README, CSS, 설정 파일까지 한 번에 수정할 수 있다.", YELLOW, False),
                ("사람의 판단", "Files changed와 Diff를 보고 요청 밖 변경이 섞였는지 확인한다.", TEAL, False),
                ("실패 비용", "잘못 merge하면 main에 불필요한 변경, 실패한 배포, secret 노출 위험이 들어간다.", NAVY, True),
            ],
        },
        {
            "label": "MAP",
            "title": "오늘은 세 개 Draft PR을 판정한다",
            "subtitle": "모든 실습 PR은 Draft라서 실제 merge하지 않습니다. 판단 연습은 Draft 배너가 아니라 파일 범위, diff, checks를 기준으로 합니다.",
            "cards": [
                ("Good PR", "README와 index처럼 예상 파일만 바뀌고 checks가 통과하는 사례", YELLOW, False),
                ("Problem A", "CI는 초록이지만 작은 copy 요청 밖 변경이 섞인 사례", TEAL, False),
                ("Problem B", ".env와 training marker, 실패한 checks로 즉시 보류할 사례", CORAL, False),
            ],
        },
        {
            "label": "GLOSSARY",
            "title": "용어는 판단 질문으로 묶어 이해한다",
            "subtitle": "단어 뜻을 작게 나열하지 않고, 실제 PR 화면에서 무엇을 확인해야 하는지 세 묶음으로 봅니다.",
            "glossary": [
                ("Issue + Branch", "Issue는 요청서, Branch는 agent가 main과 분리해 작업하는 줄기입니다.", "예: README 링크 추가 요청을 한 문장으로 줄이고 별도 branch에서 작업하게 합니다.", YELLOW),
                ("PR + Diff", "PR은 merge 전 검문소이고, Files changed/Diff는 실제 변경 증거입니다.", "예: 작은 copy 요청인데 styles.css가 보이면 scope creep을 의심합니다.", TEAL),
                ("Checks + Decision", "Checks/CI는 자동 검사이고, Merge/Hold는 사람이 내리는 운영 판단입니다.", "예: checks 실패나 .env 변경이 보이면 approve하지 않고 Hold합니다.", NAVY),
            ],
        },
        {
            "label": "FLOW",
            "title": "Issue에서 Decision까지 한 줄로 이어 읽는다",
            "subtitle": "예: README에 체크리스트 링크를 추가하라는 Issue가 있으면, PR이 정말 그 작업만 했는지 마지막 결정까지 추적합니다.",
            "flow": [
                ("Issue", "요구사항과 성공 기준을 한 문장으로 요약한다.", YELLOW),
                ("PR", "설명은 주장이다. Files changed와 Diff로 실제 변경을 확인한다.", TEAL),
                ("Decision", "Checks 결과까지 보고 Merge, Request changes, Hold 중 하나를 고른다.", NAVY),
            ],
        },
        {
            "label": "PROTOCOL",
            "title": "PR은 네 단계로 읽는다",
            "subtitle": "초보자는 Conversation부터 오래 읽기보다, 먼저 파일 목록과 실제 변경을 확인하면 agent 설명에 끌려가지 않습니다.",
            "flow": [
                ("1 Files changed", "파일 수와 종류가 Issue 범위와 맞는지 먼저 본다.", YELLOW),
                ("2 Diff", "추가/삭제 줄이 요구를 해결하는지 직접 읽는다.", WHITE),
                ("3 Checks", "CI 성공/실패와 실패 로그를 판단 근거로 쓴다.", TEAL),
                ("4 Decision", "내용상 결정을 말하되 수업 Draft PR은 실제 merge하지 않는다.", NAVY),
            ],
        },
        {
            "label": "DEMO",
            "title": "강사 데모는 Good PR 하나로 진행한다",
            "subtitle": "Good PR 하나로 Issue -> PR -> Files changed -> Checks -> Decision 클릭 흐름을 시연합니다.",
            "asset": "good-pr-overview.png",
            "links": [("Good PR #8 열기", scenarios["good-pr"]["pr_url"])],
            "bullets": [
                "Issue 링크를 열고 요구사항을 한 문장으로 줄인다",
                "Draft 배너는 안전장치로만 설명한다",
                "Files changed와 Checks 탭을 실제로 눌러 근거를 찾는다",
            ],
        },
        {
            "label": "GOOD PR",
            "title": "작은 요청, 작은 diff, 통과한 checks",
            "subtitle": "Good PR은 요청 범위 안의 파일만 바뀌고, diff가 설명과 일치하며, checks가 통과한 사례입니다.",
            "asset": "good-pr-files-changed.png",
            "chips": [("2 files", GREEN), ("Checks pass", GREEN)],
            "bullet_y": 3.55,
            "bullets": [
                "예상 파일: README.md, index.html",
                "diff가 checklist entry point 추가와 일치한다",
                "내용상 Merge 판단이지만 수업 Draft PR은 누르지 않는다",
            ],
        },
        {
            "label": "PROBLEM A",
            "title": "CI가 통과해도 scope creep이면 Hold",
            "subtitle": "Problem A는 작은 copy 요청처럼 보이지만 styles.css, README, unrelated docs가 같이 바뀐 overbroad PR입니다.",
            "asset": "problem-a-files-changed.png",
            "chips": [("4 files", RED), ("Scope creep", RED)],
            "bullet_y": 3.55,
            "bullets": [
                "Issue가 요구하지 않은 파일이 섞였는지 먼저 본다",
                "초보자는 파일 수 증가를 가장 쉬운 위험 신호로 잡는다",
                "agent에게 PR을 줄이거나 분리하라고 요청한다",
            ],
        },
        {
            "label": "PROBLEM A",
            "title": "Problem A의 Checks는 초록색이어도 충분하지 않다",
            "subtitle": "초록색 checks는 문법과 규칙을 통과했다는 뜻이지, 요청 밖 변경까지 merge해도 된다는 뜻은 아닙니다.",
            "asset": "problem-a-checks.png",
            "chips": [("Checks pass", GREEN), ("Still Hold", RED)],
            "bullet_y": 3.55,
            "bullets": [
                "CI는 자동 검사이고 scope 판단은 사람이 한다",
                "요청 밖 변경은 checks가 통과해도 보류 근거가 된다",
                "수정 요청 문장에는 '어떤 파일을 제외할지'를 적는다",
            ],
        },
        {
            "label": "PROBLEM B",
            "title": "CI 실패와 training secret marker는 Hold",
            "subtitle": f"Problem B는 .env와 {TRAINING_MARKER}가 보이고 checks도 실패하므로, 값 공유 없이 즉시 보류합니다.",
            "asset": "problem-b-files-changed.png",
            "chips": [(".env", RED), ("Do not copy value", RED)],
            "bullet_y": 3.55,
            "bullets": [
                ".env 파일은 요청과 무관해도 고위험 신호로 본다",
                "secret-like 값은 comment나 채팅에 다시 쓰지 않는다",
                "삭제뿐 아니라 history 노출 확인까지 요청한다",
            ],
        },
        {
            "label": "PROBLEM B",
            "title": "Problem B의 Checks 실패는 merge 중단 신호다",
            "subtitle": "Checks 실패는 agent가 필수 marker를 지웠거나 보안 규칙을 깼다는 신호일 수 있으므로, 실패 로그를 근거로 수정 요청합니다.",
            "asset": "problem-b-checks-failed.png",
            "chips": [("Checks fail", RED), ("Request changes", RED)],
            "bullet_y": 3.55,
            "bullets": [
                "실패한 workflow 이름과 로그의 핵심 한 줄을 찾는다",
                "값 자체가 아니라 실패 원인과 복구 요청을 적는다",
                "CI가 다시 초록색이 되기 전에는 merge 판단을 보류한다",
            ],
        },
        {
            "label": "CI/CD",
            "title": "CI는 깊게, CD는 짧게",
            "subtitle": "이번 수업의 핵심은 PR마다 돌아가는 CI를 판단 근거로 읽는 것입니다. CD는 merge 뒤 배포 자동화라는 구분만 짚습니다.",
            "cards": [
                ("CI", "PR마다 자동 검사: marker, secret-like 문자열, page title 같은 규칙을 merge 전에 확인", YELLOW, False),
                ("Checks", "GitHub PR 화면에서 workflow 성공/실패와 실패 로그를 확인하는 영역", TEAL, False),
                ("CD", "merge 뒤 자동 배포하는 흐름. 이번 실습에서는 누르거나 설정하지 않고 용어만 구분", WHITE, False),
            ],
        },
        {
            "label": "DECISION",
            "title": "Hold는 GitHub 버튼이 아니다",
            "subtitle": "Merge, Request changes, Hold는 결과만 다른 것이 아니라 필요한 근거와 GitHub에서 할 행동이 다릅니다.",
            "cards": [
                ("Merge", "요청과 diff가 일치하고 checks가 통과한다. 실제 PR이라면 합칠 수 있다.", TEAL, False),
                ("Request changes", "방향은 맞지만 수정 지점이 명확하다. 파일명과 요청을 comment로 남긴다.", YELLOW, False),
                ("Hold", "범위나 보안 위험이 커서 approve하지 않는다. 필요한 경우 blocking comment를 쓴다.", NAVY, True),
            ],
        },
        {
            "label": "COMMENT",
            "title": "좋은 agent comment는 위치, 근거, 요청을 담는다",
            "subtitle": "초보자도 바로 쓸 수 있게, comment는 '어디에서 봤는지 + 왜 문제인지 + 무엇을 고칠지' 한 문장으로 끝냅니다.",
            "bullets": [
                "위치: Files changed의 파일명, diff 줄, 또는 실패한 Checks 이름",
                "근거: Issue 범위 불일치, 불필요한 파일 변경, 실패 로그 중 하나",
                "요청: 삭제, 분리, CI 복구처럼 agent가 실행할 행동으로 끝낸다",
                "예: styles.css 변경은 요청 밖입니다. 이 PR에서 제외하거나 별도 PR로 나눠 주세요",
                "주의: secret-like 값은 comment에 복사하지 않는다",
            ],
        },
        {
            "label": "WORKTREE",
            "title": "AI agent별 작업공간을 분리하면 PR 검토가 쉬워진다",
            "subtitle": "여러 agent가 같은 폴더를 만지면 diff가 섞입니다. worktree + branch + PR로 나누면 사람은 PR 단위로만 판단하면 됩니다.",
            "cards": [
                ("기존 방식", "한 폴더에서 diff, stash, branch 전환으로 여러 agent 결과를 관리", WHITE, False),
                ("문제", "어느 agent가 어떤 파일을 바꿨는지 흐려지고, 검토 전 workspace가 지저분해짐", CORAL, False),
                ("worktree 방식", "agent마다 별도 폴더와 branch를 주고 결과를 독립 PR로 검토", TEAL, False),
            ],
        },
        {
            "label": "WORKTREE",
            "title": "Worktree는 명령 암기가 아니라 운영 패턴이다",
            "subtitle": "도구가 Codex든 Claude Code든 Antigravity든, agent별 폴더와 branch를 분리해 PR 단위로 검토한다는 원리는 같습니다.",
            "flow": [
                ("Agent A", "worktree A에서 branch A 작업 후 PR A로 제출", YELLOW),
                ("Agent B", "worktree B에서 branch B 작업 후 PR B로 제출", TEAL),
                ("Human", "PR별 Files changed, Diff, Checks를 따로 판정", NAVY),
            ],
        },
        {
            "label": "WORKTREE SETUP",
            "title": "Git worktree 생성은 세 줄로 표준화한다",
            "subtitle": "핵심은 같은 repo에서 agent별 작업 폴더와 branch를 동시에 만드는 것입니다. main을 더럽히지 않고 결과를 PR로 올립니다.",
            "command_boxes": [
                (
                    "1 최신 기준 확인",
                    ["git fetch origin", "git switch main", "git pull --ff-only"],
                    WHITE,
                    False,
                ),
                (
                    "2 agent별 worktree 생성",
                    ["git worktree add .worktrees/<tool>-<task> \\", "  -b <tool>/<task> main"],
                    YELLOW,
                    False,
                ),
                (
                    "3 해당 폴더에서 agent 실행",
                    ["cd .worktrees/<tool>-<task>", "# Codex / Claude Code / Antigravity", "# 모두 이 폴더를 workspace로 사용"],
                    NAVY,
                    True,
                ),
            ],
        },
        {
            "label": "WORKTREE SETUP",
            "title": "Codex, Claude Code, Antigravity는 이름만 다르게 쓴다",
            "subtitle": "도구마다 실행 UI는 달라도 Git 관점의 격리 방식은 같습니다. 폴더명과 branch prefix만 일관되게 정합니다.",
            "cards": [
                ("Codex", ".worktrees/codex-pr-checklist\nbranch: codex/pr-checklist", YELLOW, False),
                ("Claude Code", ".worktrees/claude-copy-fix\nbranch: claude/copy-fix", TEAL, False),
                ("Antigravity", ".worktrees/antigravity-ci-fix\nbranch: antigravity/ci-fix", NAVY, True),
            ],
        },
        {
            "label": "PRACTICE",
            "title": "개인 실습: Good PR",
            "subtitle": "혼자서 Good PR을 열고, merge 가능하다는 근거를 화면에서 직접 찾아 exit ticket 초안으로 적습니다.",
            "links": [("Good PR #8 열기", scenarios["good-pr"]["pr_url"])],
            "bullets": [
                "변경 파일 수와 파일명을 말한다",
                "diff가 Issue와 일치하는지 한 문장으로 설명한다",
                "Checks 상태가 통과인지 확인한다",
                "내용상 Merge 판단인지 말하되 실제 Draft PR은 merge하지 않는다",
            ],
        },
        {
            "label": "PRACTICE",
            "title": "페어 실습: Problem A와 Problem B",
            "subtitle": "둘이 역할을 나눠 한 명은 Files changed와 Diff를, 한 명은 Checks를 보고 왜 Hold인지 근거를 합칩니다.",
            "links": [("Problem A #9 열기", scenarios["problem-a"]["pr_url"]), ("Problem B #10 열기", scenarios["problem-b"]["pr_url"])],
            "bullets": [
                "Problem A: CI 통과와 scope creep을 분리해 판단",
                "Problem B: CI 실패와 training marker를 근거로 판단",
                "각 PR마다 agent에게 남길 수정 요청 문장을 하나씩 작성",
            ],
        },
        {
            "label": "EXIT",
            "title": "Exit ticket",
            "subtitle": "마지막 제출물은 정답 맞히기가 아니라, 판단을 화면 근거와 comment 문장으로 설명하는 연습입니다.",
            "cards": [
                ("판단", "Merge / Request changes / Hold 중 하나를 고르고 Draft PR은 실제로 누르지 않는다", YELLOW, False),
                ("근거", "Files changed, Diff, Checks 중 2개 이상을 화면 위치와 함께 쓴다", TEAL, False),
                ("comment", "위치 + 근거 + 요청이 모두 보이게 agent에게 남길 문장으로 작성", WHITE, False),
            ],
        },
    ]

    for spec in specs:
        slide = prs.slides.add_slide(blank)
        set_fill(slide)
        add_title(slide, spec["label"], spec["title"], spec["subtitle"], slide_no)
        if "cards" in spec:
            for idx, (title, body, fill, dark) in enumerate(spec["cards"]):
                add_card(slide, 0.72 + idx * 4.1, 3.25, 3.55, 1.55, title, body, fill=fill, dark=dark)
        if "command_boxes" in spec:
            for idx, (title, lines, fill, dark) in enumerate(spec["command_boxes"]):
                add_command_box(slide, 0.72 + idx * 4.1, 3.08, 3.55, 2.28, title, lines, fill=fill, dark=dark)
        if "flow" in spec:
            add_step_flow(slide, 3.35, spec["flow"])
        if "glossary" in spec:
            add_glossary_showcase(slide, spec["glossary"])
        if "bullets" in spec:
            add_bullet_stack(slide, 0.82, spec.get("bullet_y", 3.15), spec["bullets"])
        if "links" in spec:
            for idx, item in enumerate(spec["links"]):
                if isinstance(item, tuple):
                    label, url = item
                    add_link_button(slide, label, url, 0.82, 5.22 + idx * 0.46)
                else:
                    add_link_button(slide, "PR 링크 열기", item, 0.82, 5.22 + idx * 0.46)
        if "chips" in spec:
            for idx, (label, color) in enumerate(spec["chips"]):
                add_status_chip(slide, 0.82 + idx * 1.9, 3.05, label, color)
        if "asset" in spec:
            add_asset(slide, spec["asset"])
        slide_no += 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)


if __name__ == "__main__":
    create_deck()
